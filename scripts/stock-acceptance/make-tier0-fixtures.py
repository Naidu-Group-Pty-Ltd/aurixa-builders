#!/usr/bin/env python3
"""Tier-0 fixtures: one document per supported format, built from tier0_dataset.py.

    python3 scripts/stock-acceptance/make-tier0-fixtures.py . scripts/ops/fixtures/tier0

Needs reportlab, pillow, openpyxl, python-docx, python-pptx and LibreOffice
(Calc and Writer) for the legacy and OpenDocument conversions. The manifest
pins every byte the proof imports; regenerating changes the digests.

Photographs come from the acceptance corpus's own `facade()` — multi-octave
noise with the statistics of a photograph, which the product's eligibility
measure was written about — so a fixture is the kind of thing the rule judges.
Text formats carry LINK PLACEHOLDERS the proof replaces with signed links to
objects it stores in the proof organisation's own folder; binary formats carry
their photographs embedded.
"""
import csv, io, json, os, subprocess, sys, hashlib, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import importlib.util
from tier0_dataset import CSV_V1, CSV_V2, FORMATS, HEADERS, row_cells, expected

REPO = sys.argv[1]
OUT = sys.argv[2]
os.makedirs(OUT, exist_ok=True)
spec = importlib.util.spec_from_file_location('corpus', os.path.join(REPO, 'scripts/stock-acceptance/make-corpus.py'))
corpus = importlib.util.module_from_spec(spec); spec.loader.exec_module(corpus)
from PIL import Image, ImageDraw, ImageFont

def facade_bytes(seed, w=1280, h=800):
    return corpus.facade(seed, w, h).getvalue()

def plan_bytes():
    return corpus.floorplan().getvalue()

import tempfile
SCRATCH = tempfile.mkdtemp(prefix='tier0-')
def scratch_facade(seed):
    path = os.path.join(SCRATCH, f'facade-{seed}.jpg')
    if not os.path.exists(path):
        with open(path, 'wb') as fh: fh.write(facade_bytes(seed))
    return path

def write(name, data):
    path = os.path.join(OUT, name)
    with open(path, 'wb') as fh: fh.write(data)
    return path

# ---- photographs, plans and brochures (served by link) ---------------------
# Only what a LINK names is served from storage; embedded formats carry their own.
LINKED = ('tsv', 'txt', 'html', 'json', 'xml')
seeds = set()
for p in CSV_V1 + CSV_V2:
    if p.get('photo_seed'): seeds.add(p['photo_seed'])
for key in LINKED:
    for p in FORMATS[key]: seeds.add(p['photo_seed'])
os.makedirs(os.path.join(OUT, 'media'), exist_ok=True)
for s in sorted(seeds):
    write(f'media/facade-{s}.jpg', facade_bytes(s))
write('media/floorplan.jpg', plan_bytes())
# The picture a builder adds through "Add picture" where a document carries none.
write('media/remedy.jpg', facade_bytes(511))

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.lib.units import mm
W, H = A4

def brochure_pdf(p, path):
    c = canvas.Canvas(path, pagesize=A4, invariant=1)
    t = corpus.text
    t(c, 20, 28, p['design'].upper(), 20, True)
    t(c, 20, 36, f"Lot {p['lot']} {p['street']}")
    t(c, 20, 43, f"{p['estate']}, {p['suburb']} {p['state']} {p['postcode']}")
    corpus.hero(c, io.BytesIO(facade_bytes(p['photo_seed'])))
    price = f"${p['price']:,}" if not isinstance(p['price'], str) else p['price']
    t(c, 20, 192, f"Package Price - {price}", bold=True)
    t(c, 20, 199, f"Lot Size    {p['land']}m2")
    t(c, 20, 206, f"Total: {p['build']}m2")
    t(c, 20, 220, 'PROOF ONLY - not for sale. Artist impression.', 8)
    c.showPage()
    t(c, 20, 28, 'SITING PLAN', 14, True)
    t(c, 20, 40, f"Site Address: Lot {p['lot']} {p['street'].upper()}")
    t(c, 20, 47, f"Locality: {p['suburb'].upper()} ({p['postcode']})")
    t(c, 20, 54, f"State: {p['state']}")
    t(c, 20, 61, f"Home Design: {p['design'].upper()}")
    t(c, 20, 68, f"Estate: {p['estate'].upper()}")
    t(c, 20, 75, f"Site Area: {p['land']} m2")
    t(c, 20, 82, f"Build Area: {p['build']} m2")
    corpus.hero(c, corpus.floorplan(), top=200, height=110)
    t(c, 20, 210, 'This siting is subject to developer approval.', 8)
    c.showPage(); c.save()

for p in CSV_V1:
    if p.get('brochure'):
        brochure_pdf(p, os.path.join(OUT, f"media/brochure-{p['ref']}.pdf"))

# ---- the primary CSV list, v1 and v2 (links as placeholders) ---------------
LINK_HEADERS = ['Facade Image', 'Brochure', 'Floor Plan']
def link_cells(p):
    return [f"{{{{PHOTO:{p['photo_seed']}}}}}" if p.get('photo_seed') else '',
            f"{{{{BROCHURE:{p['ref']}}}}}" if p.get('brochure') else '',
            f"{{{{PLAN:{p['ref']}}}}}" if p.get('plan') else '']

def csv_bytes(props, delimiter=',', links=True):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, lineterminator='\r\n')
    w.writerow(HEADERS + (LINK_HEADERS if links else []))
    for p in props:
        w.writerow(row_cells(p) + (link_cells(p) if links else []))
    return buf.getvalue().encode('utf-8')

write('csv-v1.csv', csv_bytes(CSV_V1))
write('csv-v2.csv', csv_bytes(CSV_V2))
write('tsv.tsv', csv_bytes(FORMATS['tsv'], '\t'))
write('txt.txt', csv_bytes(FORMATS['txt'], ';'))

# ---- HTML / JSON / XML (links as placeholders) ------------------------------
import html as H_
def html_bytes(props):
    rows = ''.join('<tr>' + ''.join(f'<td>{H_.escape(c)}</td>' for c in row_cells(p) + link_cells(p)) + '</tr>' for p in props)
    head = ''.join(f'<th>{H_.escape(h)}</th>' for h in HEADERS + LINK_HEADERS)
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>Proof stock list</title></head>'
            f'<body><h1>Proof stock list (not for sale)</h1><table><thead><tr>{head}</tr></thead>'
            f'<tbody>{rows}</tbody></table></body></html>').encode('utf-8')
write('html.html', html_bytes(FORMATS['html']))

def record_dict(p):
    d = dict(zip(HEADERS + LINK_HEADERS, row_cells(p) + link_cells(p)))
    return {k: v for k, v in d.items() if v != ''}
write('json.json', json.dumps({'properties': [record_dict(p) for p in FORMATS['json']]}, ensure_ascii=False, indent=2).encode('utf-8'))
def xml_tag(h):
    import re
    return re.sub(r'[^A-Za-z0-9]+', '_', h).strip('_')
xml_rows = ''.join('<property>' + ''.join(f'<{xml_tag(k)}>{H_.escape(v)}</{xml_tag(k)}>' for k, v in record_dict(p).items()) + '</property>' for p in FORMATS['xml'])
write('xml.xml', f'<?xml version="1.0" encoding="UTF-8"?><stocklist>{xml_rows}</stocklist>'.encode('utf-8'))

# ---- XLSX with embedded, row-anchored photographs --------------------------
import openpyxl
from openpyxl.drawing.image import Image as XLImage
import datetime
def xlsx_path(props, path, title_rows=True):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Stock'
    r0 = 1
    if title_rows:
        ws.cell(row=1, column=1, value='PROOF STOCK LIST — not for sale')
        ws.cell(row=2, column=1, value='Prices include GST')
        r0 = 4
    heads = HEADERS + ['Photo']
    for j, h in enumerate(heads, 1): ws.cell(row=r0, column=j, value=h)
    for i, p in enumerate(props, 1):
        for j, v in enumerate(row_cells(p), 1): ws.cell(row=r0 + i, column=j, value=v)
        img_path = scratch_facade(p['photo_seed'])
        img = XLImage(img_path); img.width, img.height = 320, 200
        ws.add_image(img, ws.cell(row=r0 + i, column=len(heads)).coordinate)
        ws.row_dimensions[r0 + i].height = 150
    wb.properties.created = datetime.datetime(2026, 9, 28)
    wb.properties.modified = datetime.datetime(2026, 9, 28)
    wb.save(path)

for key in ('xlsx', 'xls', 'xlsm', 'ods'):
    xlsx_path(FORMATS[key], os.path.join(OUT, f'{key}.source.xlsx'))
# A workbook written by a generator that uses a DEFAULT drawing namespace
# (openpyxl and others): valid OOXML, opened by Excel, pictures row-anchored.
xlsx_path(FORMATS['xlsx-defaultns'], os.path.join(OUT, 'xlsx-defaultns.xlsx'))

# ---- DOCX with a table and embedded photographs ----------------------------
import docx
from docx.shared import Mm
def docx_path(props, path):
    d = docx.Document()
    d.add_heading('Proof stock list — not for sale', 1)
    heads = HEADERS + ['Photo']
    t = d.add_table(rows=1, cols=len(heads))
    for j, h in enumerate(heads): t.rows[0].cells[j].text = h
    for p in props:
        cells = t.add_row().cells
        for j, v in enumerate(row_cells(p)): cells[j].text = v
        cells[len(heads) - 1].paragraphs[0].add_run().add_picture(
            scratch_facade(p['photo_seed']), width=Mm(40))
    d.core_properties.created = datetime.datetime(2026, 9, 28)
    d.save(path)
for key in ('docx', 'doc', 'odt', 'rtf'):
    docx_path(FORMATS[key], os.path.join(OUT, f'{key}.source.docx'))
os.replace(os.path.join(OUT, 'docx.source.docx'), os.path.join(OUT, 'docx.docx'))

# ---- PPTX: a table slide and each property's photograph --------------------
from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation()
for p in FORMATS['pptx']:
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = f"Lot {p['lot']} {p['street']} — {p['design']}"
    tbl = s.shapes.add_table(2, len(HEADERS), Inches(0.2), Inches(1.3), Inches(9.6), Inches(0.8)).table
    for j, h in enumerate(HEADERS): tbl.cell(0, j).text = h
    for j, v in enumerate(row_cells(p)): tbl.cell(1, j).text = v
    s.shapes.add_picture(scratch_facade(p['photo_seed']), Inches(2), Inches(2.6), width=Inches(6))
prs.core_properties.created = datetime.datetime(2026, 9, 28)
prs.save(os.path.join(OUT, 'pptx.pptx'))

# ---- LibreOffice conversions: XLS, XLSM, ODS, DOC, ODT, RTF ----------------
def convert(src, target_ext, filt=None):
    outdir = os.path.join(OUT, 'lo'); os.makedirs(outdir, exist_ok=True)
    arg = f'{target_ext}:{filt}' if filt else target_ext
    subprocess.run(['soffice', '-env:UserInstallation=file:///tmp/lo-profile', '--headless', '--convert-to', arg, '--outdir', outdir, src],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    stem = os.path.splitext(os.path.basename(src))[0]
    return os.path.join(outdir, f'{stem}.{target_ext}')
for key, ext, filt in [('xls', 'xls', 'MS Excel 97'), ('xlsm', 'xlsm', 'Calc MS Excel 2007 VBA XML'),
                       ('ods', 'ods', None), ('xlsx', 'xlsx', 'Calc MS Excel 2007 XML')]:
    produced = convert(os.path.join(OUT, f'{key}.source.xlsx'), ext, filt)
    shutil.move(produced, os.path.join(OUT, f'{key}.{ext}'))
for key, ext, filt in [('doc', 'doc', 'MS Word 97'), ('odt', 'odt', None), ('rtf', 'rtf', None)]:
    produced = convert(os.path.join(OUT, f'{key}.source.docx'), ext, filt)
    shutil.move(produced, os.path.join(OUT, f'{key}.{ext}'))
for f in os.listdir(OUT):
    if '.source.' in f: os.remove(os.path.join(OUT, f))
shutil.rmtree(os.path.join(OUT, 'lo'), ignore_errors=True)

# ---- PDF schedule (a table with each row's photograph) and PDF brochure ----
def schedule_pdf(props, path):
    c = canvas.Canvas(path, pagesize=A4, invariant=1)
    corpus.text(c, 15, 20, 'KESTREL GROVE — AVAILABLE STOCK (PROOF, NOT FOR SALE)', 13, True)
    cols = [('Lot', 15), ('Address', 28), ('Design', 70), ('Bed', 98), ('Bath', 108), ('Car', 118),
            ('Land m2', 128), ('House m2', 146), ('Price', 166)]
    for label, x in cols: corpus.text(c, x, 32, label, 9, True)
    y = 40
    for p in props:
        vals = [p['lot'], p['street'], p['design'], str(p['beds']), str(p['baths']), str(p['car']),
                str(p['land']), str(p['build']), f"${p['price']:,}"]
        for (label, x), v in zip(cols, vals): corpus.text(c, x, y, v, 9)
        c.drawImage(ImageReader(io.BytesIO(facade_bytes(p['photo_seed']))), 15 * mm,
                    H - (y + 62) * mm, width=90 * mm, height=56 * mm, preserveAspectRatio=True, mask=None)
        y += 72
    c.showPage(); c.save()
schedule_pdf(FORMATS['pdf-schedule'], os.path.join(OUT, 'pdf-schedule.pdf'))
bp = FORMATS['pdf-brochure'][0]
brochure_pdf(bp, os.path.join(OUT, 'pdf-brochure.pdf'))

# ---- a PHOTOGRAPH of the schedule: PNG, JPG, WEBP, GIF ----------------------
font = ImageFont.truetype('/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf', 22)
bold = ImageFont.truetype('/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf', 24)
def schedule_image(props):
    img = Image.new('RGB', (1800, 420), (255, 255, 255)); d = ImageDraw.Draw(img)
    d.text((30, 20), 'PROOF STOCK LIST - NOT FOR SALE', font=bold, fill=(0, 0, 0))
    heads = ['Lot', 'Address', 'Suburb', 'Design', 'Bed', 'Bath', 'Car', 'Land m2', 'House m2', 'Price', 'Status']
    xs = [30, 110, 420, 600, 800, 880, 960, 1040, 1180, 1340, 1540]
    for h, x in zip(heads, xs): d.text((x, 90), h, font=bold, fill=(0, 0, 0))
    y = 150
    for p in props:
        vals = [p['lot'], p['street'], p['suburb'], p['design'], str(p['beds']), str(p['baths']),
                str(p['car']), str(p['land']), str(p['build']), f"${p['price']:,}", p['status']]
        for v, x in zip(vals, xs): d.text((x, y), v, font=font, fill=(20, 20, 20))
        y += 60
    return img
for key, fmt, ext in [('png', 'PNG', 'png'), ('jpg', 'JPEG', 'jpg'), ('webp', 'WEBP', 'webp'), ('gif', 'GIF', 'gif')]:
    buf = io.BytesIO(); schedule_image(FORMATS[key]).save(buf, format=fmt, **({'quality': 92} if fmt in ('JPEG', 'WEBP') else {}))
    write(f'{key}.{ext}', buf.getvalue())

# ---- the manifest: bytes, digests, and what each document states ----------
manifest = {'files': {}, 'expect': {}}
for root, _, files in os.walk(OUT):
    for f in sorted(files):
        if f == 'manifest.json': continue
        path = os.path.join(root, f); rel = os.path.relpath(path, OUT)
        data = open(path, 'rb').read()
        manifest['files'][rel] = {'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()}
manifest['expect']['csv-v1'] = [expected(p) for p in CSV_V1]
manifest['expect']['csv-v2'] = [expected(p) for p in CSV_V2]
for key, props in FORMATS.items():
    manifest['expect'][key] = [expected(p) for p in props]
manifest['photos'] = {p['ref']: p.get('photo_seed') for p in CSV_V1 + CSV_V2}
with open(os.path.join(OUT, 'manifest.json'), 'w') as fh: json.dump(manifest, fh, indent=1, ensure_ascii=False)
print('written', len(manifest['files']), 'files')

# ---- NEGATIVE AND EDGE INPUTS (3A) ------------------------------------------
NEG = os.path.join(OUT, 'neg'); os.makedirs(NEG, exist_ok=True)
def negw(name, data):
    with open(os.path.join(NEG, name), 'wb') as fh: fh.write(data)
H = ','.join(['Stock Ref', 'Estate', 'Lot', 'Street Address', 'Suburb', 'State', 'Postcode',
              'Design', 'Beds', 'Baths', 'Car', 'Land Size (m2)', 'House Size (m2)', 'Package Price', 'Status', 'Description'])
# unbalanced quote on the second row
negw('malformed.csv', (H + '\r\n'
     'T0-701,Kestrel Grove,701,1 Edge Road,Truganina,VIC,3029,Edge 1,3,2,2,300,150,"$500,000,Available,PROOF ONLY\r\n'
     'T0-702,Kestrel Grove,702,3 Edge Road,Truganina,VIC,3029,Edge 2,3,2,2,300,150,"$510,000",Available,PROOF ONLY\r\n').encode())
# a row missing price and counts; a row missing everything but its lot
negw('partial.csv', (H + '\r\n'
     'T0-711,Kestrel Grove,711,1 Partial Road,Truganina,VIC,3029,Part 1,,,,300,,,Available,PROOF ONLY\r\n'
     ',,712,,,,,,,,,,,,,\r\n').encode())
negw('blank.csv', b'')
negw('header-only.csv', (H + '\r\n').encode())
row721 = 'T0-721,Kestrel Grove,721,1 Twice Road,Truganina,VIC,3029,Twin 1,3,2,2,300,150,"$520,000",Available,PROOF ONLY'
negw('duplicate-rows.csv', (H + '\r\n' + row721 + '\r\n' + row721 + '\r\n'
     # same lot and estate, different ref: a second row describing the same property
     + 'T0-721B,Kestrel Grove,721,1 Twice Road,Truganina,VIC,3029,Twin 1,3,2,2,300,150,"$525,000",Available,PROOF ONLY\r\n').encode())
unusual = [
    ('T0-731', '731', '3/14 Unusual Street', '$1.2m', '4', '2.5', '2', '0.5 acres', '21.5 squares', 'Under Contract', 'Q4 2026'),
    ('T0-732', 'Lot 12A', '5 Unusual Street', '749k', '3+1', '2', '1', '420 m²', '180 sqm', 'EOI', '01/12/2026'),
    ('T0-733', 'L733', '7 Unusual Street', 'From $749,000', '3', '2', '2', '400', '200', 'Contact agent', ''),
    ('T0-734', '734', '9 Unusual Street', 'Contact agent', '99', '0', '-1', '$428,000', '12.5m x 36m', 'Sold', ''),
]
lines = [','.join(['Stock Ref', 'Estate', 'Lot', 'Street Address', 'Suburb', 'State', 'Postcode', 'Design',
                   'Beds', 'Baths', 'Car', 'Land Size (m2)', 'House Size (m2)', 'Package Price', 'Status',
                   'Completion', 'Description'])]
for ref, lot, street, price, beds, baths, car, land, build, status, comp in unusual:
    lines.append(','.join([ref, 'Kestrel Grove', f'"{lot}"', f'"{street}"', 'Truganina', 'VIC', '3029', 'Odd 1',
                           beds, baths, car, f'"{land}"', f'"{build}"', f'"{price}"', status, comp, 'PROOF ONLY']))
negw('unusual.csv', ('\r\n'.join(lines) + '\r\n').encode())
negw('unicode.csv', (H + '\r\n'
     'T0-741,Kestrel Grove — Stage ✓,741,"12 Crème Brûlée Way",Truganina,VIC,3029,"Café 25 ☕",4,2,2,448,231.5,"$749,900",Available,"PROOF ONLY — 中文说明, emoji 🏠, Ünïcödé"\r\n').encode('utf-8'))
negw('long.csv', (H + '\r\n'
     'T0-751,Kestrel Grove,751,"' + ('1 Very Long Road Name ' * 15).strip() + '",Truganina,VIC,3029,Long 1,4,2,2,448,231.5,"$749,900",Available,"'
     + ('PROOF ONLY long description. ' * 170) + '"\r\n').encode())
# not a stock list at all
Image.new('RGB', (64, 64), (120, 130, 140)).save(os.path.join(NEG, 'unsupported.tiff'), format='TIFF')
negw('program.csv', b'MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00' + b'\x00' * 2048)
import zipfile
with zipfile.ZipFile(os.path.join(NEG, 'archive.zip'), 'w') as z:
    z.writestr('stock.csv', H + '\r\nT0-761,Kestrel Grove,761,1 Zip Road,Truganina,VIC,3029,Zip 1,3,2,2,300,150,"$500,000",Available,PROOF ONLY\r\n')
# a real workbook whose name claims it is a CSV: the bytes decide
shutil.copyfile(os.path.join(OUT, 'xlsx.xlsx'), os.path.join(NEG, 'workbook-named.csv'))

manifest2 = json.load(open(os.path.join(OUT, 'manifest.json')))
for f in sorted(os.listdir(NEG)):
    data = open(os.path.join(NEG, f), 'rb').read()
    manifest2['files'][f'neg/{f}'] = {'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()}
json.dump(manifest2, open(os.path.join(OUT, 'manifest.json'), 'w'), indent=1, ensure_ascii=False)
print('negative inputs written', len(os.listdir(NEG)))
